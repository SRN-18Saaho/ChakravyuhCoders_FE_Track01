import os
import io
from datetime import date, timedelta
import numpy as np
import pandas as pd
import plotly.express as px
import streamlit as st

st.set_page_config(page_title="Mumbai Foodborne Signal Desk", page_icon="🩺", layout="wide")
ROOT=os.path.dirname(os.path.abspath(__file__)); DATA=os.path.join(ROOT,"data")
REPORT_PATH=os.path.join(DATA,"synthetic_illness_reports.csv")

@st.cache_data(ttl=60)
def load_reports():
    df=pd.read_csv(REPORT_PATH)
    for c in ["symptom_onset_date","reporting_date"]: df[c]=pd.to_datetime(df[c],errors="coerce")
    return df

REPORT_COLUMNS = [
    "report_id", "sheet_group", "ward_id", "ward_code", "ward_name", "vendor_id",
    "vendor_supplier_source", "illness_type", "severity", "symptom_onset_date",
    "reporting_date", "reporting_delay_hours", "food_type", "person_id", "batch_id",
    "hospitalised", "synthetic_data", "planted_signal_tags"
]


def load_uploaded_workbook(uploaded_file):
    """Read report rows from every usable sheet in an uploaded Excel workbook."""
    required = ["ward_name", "vendor_supplier_source", "symptom_onset_date", "reporting_date", "food_type", "severity"]
    aliases = {
        "ward": "ward_name", "ward name": "ward_name", "vendor": "vendor_supplier_source",
        "vendor / supplier source": "vendor_supplier_source", "vendor/supplier source": "vendor_supplier_source",
        "source": "vendor_supplier_source", "food type": "food_type", "illness type": "illness_type",
        "hospitalized": "hospitalised", "hospitalized?": "hospitalised", "hospitalised?": "hospitalised",
        "reporting date": "reporting_date", "symptom onset date": "symptom_onset_date",
        "ward id": "ward_id", "ward code": "ward_code", "vendor id": "vendor_id",
        "report id": "report_id", "person id": "person_id", "batch id": "batch_id",
        "reporting delay hours": "reporting_delay_hours", "synthetic data": "synthetic_data",
        "planted signal tags": "planted_signal_tags"
    }
    workbook = pd.ExcelFile(uploaded_file)
    frames = []
    skipped = []
    for sheet_name in workbook.sheet_names:
        try:
            frame = pd.read_excel(workbook, sheet_name=sheet_name)
        except Exception:
            skipped.append(sheet_name)
            continue
        if frame.empty:
            continue
        renamed = {}
        for col in frame.columns:
            raw = str(col).strip()
            key = raw.lower().replace("_", " ").strip()
            renamed[col] = aliases.get(key, raw.strip().lower().replace(" ", "_"))
        frame = frame.rename(columns=renamed)
        # Keep the expected schema when users have exported the standard Ward Reports workbook.
        if not set(required).issubset(frame.columns):
            skipped.append(sheet_name)
            continue
        frame["sheet_group"] = sheet_name
        for col in REPORT_COLUMNS:
            if col not in frame.columns:
                frame[col] = pd.NA
        frames.append(frame[REPORT_COLUMNS].copy())
    if not frames:
        raise ValueError(
            "No usable report sheets found. Each populated sheet must include ward_name, "
            "vendor_supplier_source, symptom_onset_date, reporting_date, food_type, and severity."
        )
    out = pd.concat(frames, ignore_index=True)
    # Drop entirely blank rows often left behind in spreadsheet templates.
    out = out.dropna(how="all", subset=required).copy()
    for col in ["symptom_onset_date", "reporting_date"]:
        out[col] = pd.to_datetime(out[col], errors="coerce")
    for col in ["report_id", "person_id", "ward_id", "ward_code", "vendor_id", "batch_id", "illness_type", "severity", "ward_name", "vendor_supplier_source", "food_type", "planted_signal_tags"]:
        out[col] = out[col].fillna("").astype(str).str.strip()
    out["report_id"] = [v if v else f"UPLOAD-R-{i+1:05d}" for i, v in enumerate(out["report_id"])]
    out["person_id"] = [v if v else f"UPLOAD-P-{i+1:05d}" for i, v in enumerate(out["person_id"])]
    out["illness_type"] = out["illness_type"].replace("", "Unknown / unspecified")
    out["batch_id"] = out["batch_id"].replace("nan", "").replace("None", "")
    delay = (out["reporting_date"] - out["symptom_onset_date"]).dt.total_seconds() / 3600
    out["reporting_delay_hours"] = pd.to_numeric(out["reporting_delay_hours"], errors="coerce").fillna(delay)
    def as_bool(value):
        if pd.isna(value): return False
        if isinstance(value, (bool, np.bool_)): return bool(value)
        return str(value).strip().lower() in {"yes", "y", "true", "1", "hospitalised", "hospitalized"}
    out["hospitalised"] = out["hospitalised"].map(as_bool)
    out["synthetic_data"] = out["synthetic_data"].map(lambda v: as_bool(v) if pd.notna(v) else False)
    out["planted_signal_tags"] = out["planted_signal_tags"].replace("", "NONE")
    out["source_file"] = uploaded_file.name
    return out, skipped

def daily_view(df, selected_day, include_daily_feed=True):
    """Optionally add date-seeded simulated arrivals to the current report rows."""
    seed=int(selected_day.strftime("%Y%m%d")); rng=np.random.default_rng(seed)
    # Keep historic report rows; optionally add a small date-seeded set of simulated cases.
    out=df.copy()
    n=int(rng.integers(2,10)) if include_daily_feed else 0
    wards=pd.read_csv(os.path.join(DATA,"wards_30.csv")); vendors=pd.read_csv(os.path.join(DATA,"vendors.csv")); foods=pd.read_csv(os.path.join(DATA,"food_items.csv"))
    onset_base=pd.Timestamp(selected_day)
    add=[]
    for j in range(n):
        ix=int(rng.integers(0,len(wards))); v=vendors.iloc[int(rng.integers(0,len(vendors)))]; f=foods.iloc[int(rng.integers(0,len(foods)))]
        delay=int(rng.choice([3,8,18,30,54,78],p=[.12,.2,.28,.2,.13,.07]))
        onset=onset_base-pd.Timedelta(days=int(rng.integers(0,8)))
        rep=onset+pd.Timedelta(hours=delay)
        sev=str(rng.choice(["Low","Moderate","High","Critical"],p=[.25,.46,.24,.05]))
        add.append({"report_id":f"LIVE-{selected_day:%Y%m%d}-{j+1:02d}","sheet_group":"Daily simulated feed","ward_id":wards.iloc[ix].ward_id,"ward_code":wards.iloc[ix].ward_code,"ward_name":wards.iloc[ix].ward_name,"vendor_id":v.vendor_id,"vendor_supplier_source":v.vendor_name,"illness_type":str(rng.choice(["Gastroenteritis","Acute diarrhoeal illness","Food poisoning suspected","Vomiting and diarrhoea"])),"severity":sev,"symptom_onset_date":onset,"reporting_date":rep,"reporting_delay_hours":delay,"food_type":f.food_name,"person_id":f"LIVE-P-{selected_day:%Y%m%d}-{j+1:02d}","batch_id":"","hospitalised":sev in ["High","Critical"] and bool(rng.integers(0,2)),"synthetic_data":True,"planted_signal_tags":"DAILY_SIMULATED"})
    out=pd.concat([out,pd.DataFrame(add)],ignore_index=True)
    out["reporting_date"]=pd.to_datetime(out.reporting_date,errors="coerce"); out["symptom_onset_date"]=pd.to_datetime(out.symptom_onset_date,errors="coerce")
    return out, n

def add_flags(df):
    x=df.copy(); x["red_flags"]=""
    delay=pd.to_numeric(x.reporting_delay_hours,errors="coerce")
    x.loc[delay>48,"red_flags"] += "Late report (>48h); "
    dup=x.person_id.astype(str).duplicated(keep=False)
    x.loc[dup,"red_flags"] += "Possible duplicate person ID; "
    # Isolated = no other report within same ward and 3-day onset window
    counts=[]
    for _,r in x.iterrows():
        nearby=((x.ward_id==r.ward_id)&((x.symptom_onset_date-r.symptom_onset_date).abs().dt.days<=3)).sum()-1
        counts.append(nearby)
    x["nearby_same_ward_3d"]=counts
    x.loc[x.nearby_same_ward_3d==0,"red_flags"] += "Isolated case; "
    # Popular food is a lead, not proof: flag only for investigation of potential misleading association.
    food_counts=x.groupby("food_type").report_id.transform("count")
    x.loc[(x.food_type.isin(["Vada Pav","Pav Bhaji","Samosa","Pani Puri","Biryani"])) & (food_counts>=4),"red_flags"] += "Popular-food signal (verify source); "
    # cluster heuristic: same vendor + onset window; counts do not establish causality
    for vendor, inds in x.groupby("vendor_supplier_source").groups.items():
        sub=x.loc[list(inds)]
        for idx,r in sub.iterrows():
            peers=sub[(sub.symptom_onset_date-r.symptom_onset_date).abs().dt.days<=3]
            if len(peers)>=4: x.loc[idx,"red_flags"] += "Vendor/onset cluster candidate; "
    x["red_flags"]=x.red_flags.str.strip().str.rstrip(";")
    x["flag_status"]=np.where(x.red_flags.eq(""),"No rule-based flag","Review")
    return x

st.title("🩺 Foodborne Illness Surveillance & Source Tracing")
st.caption("Mumbai / BMC ward-level investigation workspace · synthetic training data only · not a clinical diagnosis or live surveillance feed")
with st.sidebar:
    st.header("Investigation controls")
    selected_day=st.date_input("Simulation date",value=date.today())
    min_severity=st.multiselect("Severity",["Low","Moderate","High","Critical"],default=["Low","Moderate","High","Critical"])
    wards_all=pd.read_csv(os.path.join(DATA,"wards_30.csv"))
    ward_pick=st.multiselect("Ward",wards_all.ward_name.tolist())
    flags_only=st.checkbox("Show red-flagged reports only",value=False)
    st.caption("Daily synthetic arrivals are date-seeded: same date = same demo feed; next date = different arrivals.")
    st.divider()
    st.subheader("Import Excel reports")
    uploaded_excel = st.file_uploader("Upload report workbook (.xlsx)", type=["xlsx"], help="Reads populated worksheets such as Ward Reports 1–5 from the report-entry workbook.")
    import_mode = st.radio("When a workbook is uploaded", ["Combine with demo records", "Use uploaded reports instead of demo records"], index=0)
    include_daily_feed = st.checkbox("Include date-seeded demo arrivals", value=True, help="Turn this off to analyse only the saved records and any imported workbook rows.")
base=load_reports()
uploaded_count=0
if uploaded_excel is not None:
    try:
        uploaded_reports, skipped_sheets = load_uploaded_workbook(uploaded_excel)
        uploaded_count = len(uploaded_reports)
        if import_mode == "Combine with demo records":
            base = pd.concat([base, uploaded_reports], ignore_index=True, sort=False)
        else:
            base = uploaded_reports.copy()
        st.sidebar.success(f"Imported {uploaded_count:,} report(s) from Excel.")
        if skipped_sheets:
            st.sidebar.caption("Skipped sheets without the expected report columns: " + ", ".join(skipped_sheets))
    except Exception as exc:
        st.sidebar.error(f"Could not read workbook: {exc}")
data,new_n=daily_view(base,selected_day,include_daily_feed=include_daily_feed); data=add_flags(data)
# Dynamic daily KPI reflects selected reporting date, with generated daily arrivals explicitly marked
live_count=int((data.reporting_date.dt.date==selected_day).sum())
if min_severity: data=data[data.severity.isin(min_severity)]
if ward_pick: data=data[data.ward_name.isin(ward_pick)]
if flags_only: data=data[data.flag_status=="Review"]

k1,k2,k3,k4,k5=st.columns(5)
k1.metric("Reports in view",f"{len(data):,}")
k2.metric("Selected-day reports",f"{live_count:,}",help="Includes the date-seeded simulated daily feed and any matching report dates.")
k3.metric("Flagged for review",f"{(data.flag_status=='Review').sum():,}")
k4.metric("High / critical",f"{data.severity.isin(['High','Critical']).sum():,}")
k5.metric("Wards represented",f"{data.ward_id.nunique():,}")
if uploaded_excel is not None and uploaded_count:
    feed_note = f" The dashboard added {new_n} date-seeded demo report(s) for {selected_day:%d %b %Y}." if include_daily_feed else " Date-seeded demo arrivals are turned off."
    st.info(f"Excel import active: {uploaded_count:,} uploaded report(s) are included.{feed_note} Uploaded rows are treated as user-provided records; review data quality before interpreting the signals.")
else:
    if include_daily_feed:
        st.info(f"Demo feed generated {new_n} new simulated reports for {selected_day:%d %b %Y}. No real personal or patient data is used in the generated records.")
    else:
        st.info("Date-seeded demo arrivals are turned off.")

tab1,tab2,tab3,tab4,tab5,tab6,tab7=st.tabs(["Overview","Red flags & case search","Source tracing","Batch trace & common entities","Data export","Top 3 suspected sources","Enter report & Excel output"])
with tab1:
    c1,c2=st.columns(2)
    with c1:
        trend=data.groupby(data.reporting_date.dt.date).size().reset_index(name="reports"); trend.columns=["reporting_date","reports"]
        st.plotly_chart(px.line(trend,x="reporting_date",y="reports",markers=True,title="Reports by reporting date"),use_container_width=True)
    with c2:
        sev=data.groupby("severity").size().reset_index(name="reports")
        st.plotly_chart(px.bar(sev,x="severity",y="reports",color="severity",title="Severity distribution"),use_container_width=True)
    c3,c4=st.columns(2)
    with c3:
        ward_counts=data.groupby("ward_name").size().reset_index(name="reports").sort_values("reports",ascending=False).head(15)
        st.plotly_chart(px.bar(ward_counts,x="reports",y="ward_name",orientation="h",title="Reports by ward (top 15)"),use_container_width=True)
    with c4:
        foods=data.groupby("food_type").size().reset_index(name="reports").sort_values("reports",ascending=False).head(12)
        st.plotly_chart(px.bar(foods,x="food_type",y="reports",title="Frequently reported food types"),use_container_width=True)

    st.subheader("Ranked probable leads")
    st.caption(
        "Sources are ranked for investigation triage using the reports currently in view. "
        "The score is a transparent heuristic, not a probability or a finding of causation."
    )
    lead_data=data.dropna(subset=["vendor_supplier_source"]).copy()
    if lead_data.empty:
        st.info("No vendor/source values are available in the current report view to rank.")
    else:
        lead_data["person_key"]=lead_data["person_id"].astype("string").str.strip()
        lead_data.loc[lead_data["person_key"].isin(["", "<NA>", "nan", "None"]), "person_key"]=pd.NA
        lead_data["report_key"]=lead_data["report_id"].astype("string").str.strip()
        leads=lead_data.groupby("vendor_supplier_source", dropna=True).agg(
            reports=("report_key", "nunique"),
            unique_people=("person_key", "nunique"),
            wards_affected=("ward_id", "nunique"),
            high_critical=("severity", lambda s: s.isin(["High", "Critical"]).sum()),
            first_onset=("symptom_onset_date", "min"),
            last_onset=("symptom_onset_date", "max"),
        ).reset_index()
        # Normalize each dimension against the strongest source in this current view.
        # Weights are intentionally visible and illustrative, not epidemiologically validated.
        score=pd.Series(0.0, index=leads.index)
        for column, weight in [("unique_people", 0.40), ("wards_affected", 0.25), ("high_critical", 0.25), ("reports", 0.10)]:
            maximum=leads[column].max()
            if pd.notna(maximum) and maximum > 0:
                score += leads[column] / maximum * weight
        leads["priority_score"]=(score*100).round(1)
        leads=leads.sort_values(["priority_score", "unique_people", "high_critical", "reports"], ascending=False).reset_index(drop=True)
        leads.insert(0, "rank", range(1, len(leads)+1))
        leads=leads.rename(columns={
            "vendor_supplier_source":"Probable lead / source",
            "reports":"Reports",
            "unique_people":"Unique person IDs",
            "wards_affected":"Wards affected",
            "high_critical":"High / critical reports",
            "first_onset":"First symptom onset",
            "last_onset":"Latest symptom onset",
            "priority_score":"Priority score / 100",
        })
        for col in ["First symptom onset", "Latest symptom onset"]:
            leads[col]=pd.to_datetime(leads[col], errors="coerce").dt.strftime("%Y-%m-%d").fillna("Unknown")
        st.dataframe(leads, use_container_width=True, hide_index=True)
        st.caption(
            "Scoring weights: unique people 40%, wards affected 25%, high/critical reports 25%, "
            "and distinct report IDs 10%, each normalized against the highest source value in the current view. "
            "Repeated person IDs do not inflate the unique-person component. Popular foods, delayed reports, "
            "isolated cases, and missing exposure details can mislead; review the red flags and verify exposure, "
            "batch, inspection, and laboratory evidence before drawing conclusions."
        )
    st.subheader("Ward coverage")
    coverage=wards_all.merge(data.groupby("ward_id").size().rename("reports"),on="ward_id",how="left").fillna({"reports":0})
    st.dataframe(coverage.rename(columns={"ward_name":"Ward","ward_code":"Ward code","zone":"Zone","reports":"Reports"})[["Ward","Ward code","Zone","Reports"]],use_container_width=True,hide_index=True)
with tab2:
    st.subheader("Search by person ID, vendor, report ID, food, or ward")
    q=st.text_input("Search records",placeholder="e.g. SYN-P-0001, Anand Snacks, Vada Pav")
    view=data.copy()
    if q:
        mask=view.astype(str).apply(lambda col: col.str.contains(q,case=False,na=False)).any(axis=1); view=view[mask]
    cols=["report_id","person_id","ward_name","vendor_supplier_source","illness_type","severity","symptom_onset_date","reporting_date","reporting_delay_hours","food_type","red_flags"]
    st.dataframe(view[cols].sort_values("reporting_date",ascending=False),use_container_width=True,hide_index=True)
    st.caption("Rules are triage heuristics. A popular food alone is not evidence of contamination; isolated cases are not automatically low-risk. Verify records and interview/exposure data.")
with tab3:
    st.subheader("Vendor / supplier source concentration")
    source=(data.groupby(["vendor_supplier_source","food_type"], dropna=False)
        .agg(reports=("report_id","nunique"),
             ward_count=("ward_id","nunique"),
             high_critical=("severity",lambda s:s.isin(["High","Critical"]).sum()),
             first_onset=("symptom_onset_date","min"),
             last_onset=("symptom_onset_date","max"))
        .reset_index())
    # Show ward names and codes together; preserve distinct wards and sort for easy scanning.
    ward_labels=(data.dropna(subset=["ward_id"])
        .assign(ward_label=lambda d: d["ward_name"].fillna("Unknown ward").astype(str)
                + " (" + d["ward_code"].fillna("No code").astype(str) + ")")
        .groupby(["vendor_supplier_source","food_type"], dropna=False)["ward_label"]
        .agg(lambda values: ", ".join(sorted(set(values))))
        .rename("wards"))
    source=source.merge(ward_labels.reset_index(), on=["vendor_supplier_source","food_type"], how="left")
    source=source.sort_values(["reports","high_critical","ward_count"],ascending=False).reset_index(drop=True)
    source.insert(0, "Rank", range(1, len(source)+1))
    source=source.rename(columns={"vendor_supplier_source":"Vendor / Supplier Source",
                                  "food_type":"Food Type", "reports":"Reports",
                                  "ward_count":"Number of wards", "wards":"Ward name (ward code)",
                                  "high_critical":"High / critical reports",
                                  "first_onset":"First symptom onset", "last_onset":"Latest symptom onset"})
    st.dataframe(source,use_container_width=True,hide_index=True)
    st.subheader("Potential source investigation")
    source_name=st.selectbox("Select source/vendor",sorted(data.vendor_supplier_source.dropna().unique()))
    sub=data[data.vendor_supplier_source==source_name]
    st.write(f"**{source_name}** — {len(sub)} reports across {sub.ward_name.nunique()} wards; {sub.severity.isin(['High','Critical']).sum()} high/critical reports.")
    st.dataframe(sub[["report_id","ward_name","person_id","food_type","batch_id","symptom_onset_date","reporting_date","severity","red_flags"]],use_container_width=True,hide_index=True)
    st.warning("Source ranking is a prioritization aid, not causal attribution. Confirm with exposure histories, batch records, inspection results, and lab evidence.")

with tab4:
    st.subheader("Recorded supply-chain batch tracing")
    st.caption("This view uses the supplied supply_chain_events.csv. Links show recorded movements, not proof that a batch caused illness.")
    chain_path=os.path.join(DATA,"supply_chain_events.csv")
    if os.path.exists(chain_path):
        chain=pd.read_csv(chain_path)
        needed={"batch_id","from_entity_id","from_entity_type","to_entity_id","to_entity_type","event_type"}
        if needed.issubset(chain.columns):
            available_batches=sorted(set(data.batch_id.dropna().astype(str)) & set(chain.batch_id.dropna().astype(str)))
            if not available_batches:
                available_batches=sorted(chain.batch_id.dropna().astype(str).unique())
                st.info("No batch IDs overlap between the current report view and supply-chain records. Showing recorded supply-chain batches for exploration.")
            if available_batches:
                selected_batch=st.selectbox("Select batch",available_batches,key="trace_batch")
                edges=chain[chain.batch_id.astype(str)==selected_batch].copy()
                st.dataframe(edges,use_container_width=True,hide_index=True)
                entities=sorted(set(edges.from_entity_id.dropna().astype(str)) | set(edges.to_entity_id.dropna().astype(str)))
                if entities:
                    entity=st.selectbox("Trace from entity",entities,key="trace_entity")
                    upstream_map={}
                    downstream_map={}
                    for row in edges.itertuples(index=False):
                        src=str(row.from_entity_id); dst=str(row.to_entity_id)
                        upstream_map.setdefault(dst,set()).add(src)
                        downstream_map.setdefault(src,set()).add(dst)
                    def walk_trace(start, adjacency):
                        seen=set(); stack=list(adjacency.get(start,set()))
                        while stack:
                            node=stack.pop()
                            if node==start or node in seen: continue
                            seen.add(node); stack.extend(adjacency.get(node,set())-seen)
                        return sorted(seen)
                    upstream=walk_trace(entity,upstream_map); downstream=walk_trace(entity,downstream_map)
                    a,b=st.columns(2)
                    with a:
                        st.markdown(f"**Upstream of {entity}**")
                        st.write(" → ".join(upstream) if upstream else "No recorded upstream links.")
                    with b:
                        st.markdown(f"**Downstream of {entity}**")
                        st.write(" → ".join(downstream) if downstream else "No recorded downstream links.")
            common_rows=[]
            for row in chain.dropna(subset=["batch_id","from_entity_id","to_entity_id"]).itertuples(index=False):
                common_rows.extend([{"entity":str(row.from_entity_id),"batch_id":str(row.batch_id)},
                                    {"entity":str(row.to_entity_id),"batch_id":str(row.batch_id)}])
            if common_rows:
                common=pd.DataFrame(common_rows).drop_duplicates().groupby("entity").agg(
                    batch_count=("batch_id","nunique"),
                    batch_ids=("batch_id",lambda x:", ".join(sorted(x.unique())))
                ).reset_index()
                common=common[common.batch_count>1].sort_values(["batch_count","entity"],ascending=[False,True])
                st.subheader("Entities shared across multiple batches")
                if common.empty: st.info("No entity is recorded across multiple batches.")
                else: st.dataframe(common,use_container_width=True,hide_index=True)
        else:
            st.error("Supply-chain CSV is missing required columns.")
    else:
        st.info("No supply_chain_events.csv found in the data folder.")


with tab6:
    st.subheader("Top 3 suspected sources")
    st.caption(
        "A focused shortlist of the three highest-ranked vendor/supplier leads in the current filtered report view. "
        "Rankings are triage heuristics only—not probabilities, diagnoses, or proof of contamination."
    )
    top_data=data.dropna(subset=["vendor_supplier_source"]).copy()
    if top_data.empty:
        st.info("No vendor/supplier source values are available in the current report view to rank.")
    else:
        top_data["person_key"]=top_data["person_id"].astype("string").str.strip()
        top_data.loc[top_data["person_key"].isin(["", "<NA>", "nan", "None"]), "person_key"]=pd.NA
        top_data["report_key"]=top_data["report_id"].astype("string").str.strip()
        top_leads=top_data.groupby("vendor_supplier_source", dropna=True).agg(
            reports=("report_key", "nunique"),
            unique_people=("person_key", "nunique"),
            wards_affected=("ward_id", "nunique"),
            high_critical=("severity", lambda s: s.isin(["High", "Critical"]).sum()),
            first_onset=("symptom_onset_date", "min"),
            last_onset=("symptom_onset_date", "max"),
            food_types=("food_type", lambda s: ", ".join(sorted(set(s.dropna().astype(str))))),
        ).reset_index()
        score=pd.Series(0.0, index=top_leads.index)
        for column, weight in [("unique_people", 0.40), ("wards_affected", 0.25), ("high_critical", 0.25), ("reports", 0.10)]:
            maximum=top_leads[column].max()
            if pd.notna(maximum) and maximum > 0:
                score += top_leads[column] / maximum * weight
        top_leads["priority_score"]=(score*100).round(1)
        top_leads=top_leads.sort_values(
            ["priority_score", "unique_people", "high_critical", "reports"],
            ascending=False
        ).head(3).reset_index(drop=True)
        top_leads.insert(0, "Rank", range(1, len(top_leads)+1))
        top_leads=top_leads.rename(columns={
            "vendor_supplier_source":"Suspected source",
            "reports":"Distinct reports",
            "unique_people":"Unique people",
            "wards_affected":"Wards affected",
            "high_critical":"High / critical reports",
            "first_onset":"First symptom onset",
            "last_onset":"Latest symptom onset",
            "food_types":"Reported food types",
            "priority_score":"Priority score / 100",
        })
        for col in ["First symptom onset", "Latest symptom onset"]:
            top_leads[col]=pd.to_datetime(top_leads[col], errors="coerce").dt.strftime("%Y-%m-%d").fillna("Unknown")
        st.dataframe(top_leads, use_container_width=True, hide_index=True)
        st.markdown("**How to interpret the shortlist**")
        st.write(
            "Scores combine unique people (40%), wards affected (25%), high/critical reports (25%), "
            "and distinct report IDs (10%). Each measure is normalized against the highest source value "
            "in the current filtered view. A popular food, a shared vendor, late reporting, duplicate IDs, "
            "or an isolated case can create misleading signals. Confirm with exposure interviews, batch records, "
            "inspections, and laboratory evidence before drawing conclusions."
        )
        st.download_button(
            "Download Top 3 shortlist as CSV",
            data=top_leads.to_csv(index=False).encode("utf-8"),
            file_name=f"top_3_suspected_sources_{selected_day.isoformat()}.csv",
            mime="text/csv",
        )


with tab7:
    st.subheader("Enter a new illness report")
    st.caption("Enter the seven requested fields. Ward/vendor identifiers, report/person IDs, reporting delay, and scenario tags are calculated automatically. Records entered here are held in this browser session until it is refreshed or closed.")
    wards_input=pd.read_csv(os.path.join(DATA,"wards_30.csv"))
    vendors_input=pd.read_csv(os.path.join(DATA,"vendors.csv"))
    foods_input=pd.read_csv(os.path.join(DATA,"food_items.csv"))
    illness_options=sorted(set(base.illness_type.dropna().astype(str)) | {"Acute diarrhoeal illness","Food poisoning suspected","Gastroenteritis","Nausea / vomiting","Vomiting and diarrhoea","Other / unspecified"})
    severity_options=["Low","Moderate","High","Critical"]
    food_options=sorted(set(foods_input.food_name.dropna().astype(str)) | set(base.food_type.dropna().astype(str)))
    vendor_options=sorted(set(vendors_input.vendor_name.dropna().astype(str)) | set(base.vendor_supplier_source.dropna().astype(str)))
    with st.form("new_illness_report_form", clear_on_submit=True):
        f1,f2=st.columns(2)
        with f1:
            entered_ward=st.selectbox("1. Ward name",wards_input.ward_name.tolist())
            entered_vendor_choice=st.selectbox("2. Vendor / supplier source",vendor_options+["Other / enter manually"])
            entered_vendor_custom=st.text_input("Enter vendor / supplier source",placeholder="Required for a new source") if entered_vendor_choice=="Other / enter manually" else ""
            entered_onset=st.date_input("3. Symptom onset date",value=date.today())
            entered_reporting=st.date_input("4. Reporting date",value=date.today())
        with f2:
            entered_food_choice=st.selectbox("5. Food type",food_options+["Other / enter manually"])
            entered_food_custom=st.text_input("Enter food type",placeholder="Required for a new food type") if entered_food_choice=="Other / enter manually" else ""
            entered_hospitalised=st.selectbox("6. Hospitalized?",["No","Yes"])
            entered_illness=st.selectbox("7a. Illness type",illness_options)
            entered_severity=st.selectbox("7b. Severity",severity_options,index=1)
        add_report=st.form_submit_button("Add report to output",type="primary")
    if "user_entered_reports" not in st.session_state:
        st.session_state.user_entered_reports=[]
    if add_report:
        entered_vendor=entered_vendor_custom.strip() if entered_vendor_choice=="Other / enter manually" else entered_vendor_choice
        entered_food=entered_food_custom.strip() if entered_food_choice=="Other / enter manually" else entered_food_choice
        if not entered_vendor or not entered_food:
            st.error("Please enter both the new vendor/source and food type when selecting ‘Other / enter manually’.")
        elif entered_reporting < entered_onset:
            st.error("Reporting date cannot be earlier than symptom onset date. Please correct the dates and submit again.")
        else:
            wardrow=wards_input.loc[wards_input.ward_name==entered_ward].iloc[0]
            vendorrow=vendors_input.loc[vendors_input.vendor_name==entered_vendor]
            vendor_id=str(vendorrow.iloc[0].vendor_id) if not vendorrow.empty else ""
            # Use existing report numbering plus the number of reports entered this session.
            old_report_nums=pd.to_numeric(base.report_id.astype(str).str.extract(r"(\d+)$")[0],errors="coerce")
            next_report=int(old_report_nums.max() or 0)+len(st.session_state.user_entered_reports)+1
            person_nums=pd.to_numeric(base.person_id.astype(str).str.extract(r"(\d+)$")[0],errors="coerce")
            next_person=int(person_nums.max() or 0)+len(st.session_state.user_entered_reports)+1
            delay=int((pd.Timestamp(entered_reporting)-pd.Timestamp(entered_onset)).total_seconds()/3600)
            tags=[]
            if delay>48: tags.append("LATE_REPORT")
            if entered_food in ["Vada Pav","Pav Bhaji","Samosa","Pani Puri","Biryani"]: tags.append("POPULAR_FOOD_REVIEW")
            # Check for possible same-ward onset cluster in the existing and session-entered reports.
            other=pd.concat([base, pd.DataFrame(st.session_state.user_entered_reports)],ignore_index=True) if st.session_state.user_entered_reports else base
            other_onset=pd.to_datetime(other.symptom_onset_date,errors="coerce")
            peer=((other.ward_id.astype(str)==str(wardrow.ward_id)) & ((other_onset-pd.Timestamp(entered_onset)).abs().dt.days<=3)).sum()
            if peer==0: tags.append("ISOLATED_CASE_REVIEW")
            row={
                "report_id":f"USR-R-{next_report:05d}","sheet_group":"User-entered report",
                "ward_id":str(wardrow.ward_id),"ward_code":str(wardrow.ward_code),"ward_name":str(wardrow.ward_name),
                "vendor_id":vendor_id,"vendor_supplier_source":entered_vendor,"illness_type":entered_illness,
                "severity":entered_severity,"symptom_onset_date":entered_onset.isoformat(),"reporting_date":entered_reporting.isoformat(),
                "reporting_delay_hours":delay,"food_type":entered_food,"person_id":f"USR-P-{next_person:05d}",
                "batch_id":"","hospitalised":entered_hospitalised=="Yes","synthetic_data":False,
                "planted_signal_tags":";".join(tags) if tags else "NONE"
            }
            st.session_state.user_entered_reports.append(row)
            st.success(f"Added report {row['report_id']}. Reporting delay: {delay} hours.")
    entered_rows=st.session_state.user_entered_reports
    output_columns=["report_id","ward_id","ward_code","ward_name","vendor_id","vendor_supplier_source","illness_type","severity","symptom_onset_date","reporting_date","reporting_delay_hours","food_type","person_id","batch_id","hospitalised","synthetic_data","planted_signal_tags"]
    if entered_rows:
        entered_df=pd.DataFrame(entered_rows)
        st.markdown("**Reports queued for Excel output**")
        st.dataframe(entered_df[output_columns],use_container_width=True,hide_index=True)
        # Also show rule-based review notes without changing the established 17-column workbook layout.
        st.markdown("**Calculated review checks**")
        review=entered_df[["report_id","reporting_delay_hours","planted_signal_tags"]].copy()
        review["review_note"]=review.planted_signal_tags.apply(lambda x: "; ".join({"LATE_REPORT":"Late report (>48 hours)","POPULAR_FOOD_REVIEW":"Popular-food signal; verify exposure/source","ISOLATED_CASE_REVIEW":"No similar same-ward onset in prior records"}.get(t,t) for t in str(x).split(";") if t and t!="NONE") or "No automatic review tag")
        st.dataframe(review.drop(columns=["planted_signal_tags"]),use_container_width=True,hide_index=True)
        excel_buffer=io.BytesIO()
        with pd.ExcelWriter(excel_buffer,engine="openpyxl") as writer:
            # Keep the same five sheet names and the same 17-column schema as the previous workbook.
            for i in range(5):
                chunk=entered_df.iloc[i*40:(i+1)*40]
                chunk[output_columns].to_excel(writer,index=False,sheet_name=f"Ward Reports {i+1}")
                ws=writer.sheets[f"Ward Reports {i+1}"]
                ws.freeze_panes="A2"
                ws.auto_filter.ref=ws.dimensions
                from openpyxl.styles import Font, PatternFill, Alignment
                for cell in ws[1]:
                    cell.font=Font(bold=True,color="FFFFFF")
                    cell.fill=PatternFill("solid",fgColor="17365D")
                    cell.alignment=Alignment(wrap_text=True,vertical="center")
                ws.row_dimensions[1].height=30
                for column_cells in ws.columns:
                    letter=column_cells[0].column_letter
                    max_len=max((len(str(c.value)) if c.value is not None else 0) for c in column_cells[:100])
                    ws.column_dimensions[letter].width=min(max(max_len+2,12),34)
        st.download_button("Download entered reports as Excel workbook",data=excel_buffer.getvalue(),file_name=f"user_entered_surveillance_reports_{date.today().isoformat()}.xlsx",mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",type="primary")
        st.download_button("Download entered reports as CSV",data=entered_df[output_columns].to_csv(index=False).encode("utf-8"),file_name=f"user_entered_surveillance_reports_{date.today().isoformat()}.csv",mime="text/csv")
        if st.button("Clear entered reports",key="clear_user_reports"):
            st.session_state.user_entered_reports=[]
            st.rerun()
        st.caption(f"{len(entered_df)} report(s) queued. Workbook retains the previous five sheet names and exact 17-column field order; up to 40 entered reports go on each sheet. Unused sheets remain empty.")
    else:
        st.info("No reports entered yet. Complete the form and select ‘Add report to output’ to begin.")

with tab5:
    st.subheader("Export filtered report view")
    st.download_button("Download CSV",data.to_csv(index=False).encode("utf-8"),file_name=f"surveillance_{selected_day.isoformat()}.csv",mime="text/csv")
    st.download_button("Download five-sheet synthetic workbook",open(os.path.join(ROOT,"synthetic_surveillance_reports.xlsx"),"rb").read(),file_name="synthetic_surveillance_reports.xlsx",mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    st.subheader("Data dictionary")
    st.markdown("- **person_id**: synthetic deduplication key; duplicates are deliberately planted for testing.\n- **vendor_supplier_source**: candidate exposure source, not proven source.\n- **reporting_delay_hours**: onset-to-report delay; >48 hours flagged.\n- **planted_signal_tags**: scenario annotations for QA; should not be treated as field evidence.\n- **synthetic_data**: true for all generated demo rows.")
